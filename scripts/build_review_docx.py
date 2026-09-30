#!/usr/bin/env python3
"""One Word file with all the embedding-analysis material, for the author to read and comment.

    /home/vasco/anaconda3/bin/python scripts/build_review_docx.py

User, 2026-09-30: nine separate files are too many, put them together. Each source becomes one
Part: its own title is replaced by the Part heading, its sections keep their levels, and every
Part starts on a new page after a table of contents (Word fills it on opening).

The reports are generated Markdown with HTML figure blocks (the PDF route renders the HTML);
Word needs plain Markdown images, so each <figure> becomes ``![caption](<absolute src>)``, the
<style> blocks and wrapper <div>s go, and inline tags become Markdown. Image paths are made
absolute because the Parts come from different folders. Tables are already pipe tables; they
get grid borders after pandoc (its default table style has none).

Output: ``3 - Images DataSets & Labelling Outputs/Embedding Analysis brief/current/``. The pack
shows © evtol.news photos, so the whole file is private. Commented copies go to ``../comments/``
(the Labelling Analysis brief convention).
"""

from __future__ import annotations

import html
import re
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import quote

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

REPO = Path(__file__).resolve().parent.parent
DATA = Path("/mnt/storage_11tb/Drive_files_to_syncronize/3 - Images DataSets & Labelling Outputs")
PAT, EVN = DATA / "1639_LABELLED/3_embedding_evaluation", DATA / "EVTOLNEWS_DS/3_embedding_evaluation"
DOCS = REPO / "docs/embedding_evaluation"
CURRENT = DATA / "Embedding Analysis brief/current"
#: user, 2026-09-30: 82 pages is too much. The summary is what to read; the pack is the reference.
SUMMARY_OUT = CURRENT / "1 Embedding Analysis — summary (read this).docx"
PACK_OUT = CURRENT / "2 Embedding Analysis — full reference (private).docx"

#: (Part title, source markdown, extra images appended as figures: [(path, caption)])
PARTS = [
    ("How to read this pack", DOCS / "REVIEW_GUIDE.md", []),
    ("Chapter plan and decisions", DOCS / "CHAPTER_PLAN.md", []),
    ("Test: the registers model and its attention", PAT / "registers_check/REGISTERS_CHECK.md", []),
    ("Test: view erasure", PAT / "view_erasure/VIEW_ERASURE.md", []),
    ("Test: photo crop", EVN / "crop_experiment/RESULTS.md",
     [(EVN / "crop_experiment/montage_orig_vs_crop.png",
       "Original (left) vs OWLv2 crop (right), 40 stratified pairs, 8 per class; label under each pair: "
       "class, aircraft, detector score, crop share of the image.")]),
    ("Patent report: base model", DOCS / "PATENT_EMBEDDING_ANALYSIS.md", []),
    ("Patent report: registers model", DOCS / "PATENT_EMBEDDING_ANALYSIS_REGISTERS.md", []),
    ("Photo report: base model", EVN / "embedding_analysis/EVTOLNEWS_EMBEDDING_ANALYSIS.md", []),
    ("Photo report: registers model",
     EVN / "embedding_analysis_registers/EVTOLNEWS_EMBEDDING_ANALYSIS_REGISTERS.md", []),
]

#: rotatable 3-D pages, listed in the last Part as file links
PAGES_3D = [
    ("Patent figures, base model", PAT / "umap_3d/PATENT_UMAP_3D.html",
     "12-class kNN 0.18, 5-class 0.33"),
    ("Patent figures, registers model", PAT / "umap_3d/PATENT_UMAP_3D_REGISTERS.html",
     "12-class kNN 0.235, 5-class 0.445"),
    ("evtol.news photos (with the patents in one joint map)",
     EVN / "advisor_report/EVTOLNEWS_UMAP_3D.html", "photos 5-class kNN 0.53"),
]

PAGE_BREAK = '\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'
FIGURE = re.compile(r'<figure[^>]*>\s*<img src="([^"]+)"(?:[^>]*style="width:(\d+)%")?[^>]*>'
                    r'\s*(?:<figcaption>(.*?)</figcaption>)?\s*</figure>', re.S)
MD_IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)\s<]+)\)(\{width=(\d+)%\})?")


def _inline(s: str) -> str:
    s = re.sub(r"<(strong|b)>(.*?)</\1>", r"**\2**", s, flags=re.S)
    s = re.sub(r"<(em|i)>(.*?)</\1>", r"*\2*", s, flags=re.S)
    s = re.sub(r"<code>(.*?)</code>", r"`\1`", s, flags=re.S)
    s = re.sub(r'<span class="src">', " ", s)
    s = re.sub(r"<br\s*/?>", " ", s)
    s = re.sub(r"</?(span|a|sup|sub|small|p)[^>]*>", "", s)
    return html.unescape(s)


def _image(src: str, base: Path, cap: str, width: str = "100") -> str:
    path = src if src.startswith("/") else str((base / src).resolve())
    cap = re.sub(r"\s+", " ", cap).strip().replace("[", "(").replace("]", ")")
    return f"\n\n![{cap}](<{path}>){{width={width}%}}\n\n"


def to_markdown(text: str, base: Path) -> str:
    text = re.sub(r"<style>.*?</style>", "", text, flags=re.S)
    text = FIGURE.sub(lambda m: _image(m.group(1), base, _inline(m.group(3) or ""), m.group(2) or "100"), text)
    text = MD_IMAGE.sub(lambda m: _image(m.group(2), base, m.group(1), m.group(4) or "100"), text)
    text = re.sub(r"^\s*</?div[^>]*>\s*$", "", text, flags=re.M)
    text = re.sub(r"^# .*$", "", text, count=1, flags=re.M)   # the source title gives way to the Part
    return _inline(text)


def links_part() -> str:
    rows = ["These pages open in the browser and can be rotated; each one carries a table with the "
            "kNN check, so the picture can be judged against a number. They are single files: to "
            "show them on another computer, send the file itself.", ""]
    for name, path, nums in PAGES_3D:
        url = "file://" + quote(str(path))
        rows += [f"- **{name}** ({nums}): [{path.name}]({url})", f"  `{path}`"]
    return "\n".join(rows)


def _borders(doc: Document) -> None:
    for t in doc.tables:
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "BFBFBF")
            b.append(e)
        t._tbl.tblPr.append(b)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8.5)


def _write(md: str, out: Path, title: str, subtitle: str, toc: bool) -> Path:
    head = f'---\ntitle: "{title}"\nsubtitle: "{subtitle}"\ndate: 2026-09-30\n---\n\n'
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as f:
        f.write(head + md)
    subprocess.run(["pandoc", f.name, "-f", "markdown", "-t", "docx", "-o", str(out)]
                   + (["--toc", "--toc-depth=2"] if toc else []), check=True)
    Path(f.name).unlink()
    doc = Document(out)
    _borders(doc)
    if toc:
        upd = OxmlElement("w:updateFields")   # pandoc writes the TOC field empty; Word fills it on opening
        upd.set(qn("w:val"), "true")
        doc.settings.element.append(upd)
    doc.save(out)
    return out


def build_summary() -> Path:
    src = DOCS / "REVIEW_SUMMARY.md"
    return _write(to_markdown(src.read_text(encoding="utf-8"), src.parent), SUMMARY_OUT,
                  "Embedding Analysis: summary", "DINOv2 chapter, the model decision and what the chapter "
                  "can claim (private: one figure shows © photos)", toc=False)


def build_pack() -> Path:
    parts = []
    for i, (title, src, extra) in enumerate(PARTS):
        md = to_markdown(src.read_text(encoding="utf-8"), src.parent)
        for img, cap in extra:
            md += _image(str(img), src.parent, cap)
        parts.append(f"# Part {i}. {title}\n\n{md}")
    parts.append(f"# Part {len(PARTS)}. The 3-D views\n\n{links_part()}")
    return _write(PAGE_BREAK.join(parts), PACK_OUT, "Embedding Analysis: full reference",
                  "Every report and test behind the summary (private: shows © photos)", toc=True)


if __name__ == "__main__":
    import sys
    targets = sys.argv[1:] or ["summary", "pack"]
    for t in targets:
        print({"summary": build_summary, "pack": build_pack}[t]())
